from pathlib import Path
from math import cos, sin, atan2, pi

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("documentacion") / "Administracion de proyectos - propuesta documental sin notas.docx"
ASSETS = Path("tools") / "doc_assets"
BANNER = ASSETS / "banner_utvm.jpg"
LOGO = ASSETS / "logo_utvm_recorte.png"
USE_CASE_IMG = ASSETS / "modelo_casos_uso.png"
ER_IMG = ASSETS / "modelo_entidad_relacion.png"
REL_IMG = ASSETS / "modelo_relacional_completo.png"
SITE_IMG = ASSETS / "mapa_sitio.png"

BLUE = "15529A"
BLUE_DARK = "0D376D"
GREEN = "21A366"
GREEN_DARK = "0F7D47"
CYAN = "1EA7D7"
GOLD = "F3B21A"
INK = "1F2937"
MUTED = "5B677A"
LIGHT_BLUE = "EAF3FB"
LIGHT_GREEN = "EAF7EF"
LIGHT_GRAY = "F5F7FA"
BORDER = "CAD6E2"


def fnt(size=24, bold=False):
    candidates = [
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf",
    ]
    for c in candidates:
        if Path(c).exists():
            return ImageFont.truetype(c, size)
    return ImageFont.load_default()


def rgb(hex_color):
    return tuple(int(hex_color[i:i + 2], 16) for i in (0, 2, 4))


def wrap(draw, text, font, max_width):
    words = text.split()
    lines, current = [], ""
    for word in words:
        test = word if not current else current + " " + word
        if draw.textbbox((0, 0), test, font=font)[2] <= max_width:
            current = test
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def center(draw, box, text, font, fill=rgb(INK), max_width=None):
    x1, y1, x2, y2 = box
    lines = wrap(draw, text, font, max_width or x2 - x1 - 20)
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total = sum(heights) + 5 * (len(lines) - 1)
    y = y1 + (y2 - y1 - total) / 2
    for line, h in zip(lines, heights):
        bb = draw.textbbox((0, 0), line, font=font)
        x = x1 + (x2 - x1 - (bb[2] - bb[0])) / 2
        draw.text((x, y), line, font=font, fill=fill)
        y += h + 5


def arrow(draw, start, end, fill=rgb(MUTED), width=3):
    draw.line([start, end], fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    theta = atan2(y2 - y1, x2 - x1)
    angle = pi / 7
    length = 15
    p1 = (x2 - length * cos(theta - angle), y2 - length * sin(theta - angle))
    p2 = (x2 - length * cos(theta + angle), y2 - length * sin(theta + angle))
    draw.polygon([end, p1, p2], fill=fill)


def poly_arrow(draw, points, fill=rgb(MUTED), width=3):
    for i in range(len(points) - 1):
        draw.line([points[i], points[i + 1]], fill=fill, width=width)
    arrow(draw, points[-2], points[-1], fill=fill, width=width)


def make_logo():
    if not BANNER.exists():
        return
    img = Image.open(BANNER).convert("RGB")
    w, h = img.size
    crop = img.crop((0, 0, min(w, 460), min(h, 210)))
    crop.save(LOGO)


def entity(draw, x, y, w, title, attrs, fill):
    row_h = 30
    h = 46 + row_h * len(attrs)
    draw.rounded_rectangle((x, y, x + w, y + h), radius=8, fill="white", outline=rgb(BLUE_DARK), width=3)
    draw.rectangle((x, y, x + w, y + 44), fill=rgb(fill), outline=rgb(BLUE_DARK), width=3)
    center(draw, (x, y, x + w, y + 44), title, fnt(20, True), fill=rgb(INK))
    yy = y + 54
    for a in attrs:
        draw.text((x + 12, yy), a, font=fnt(17), fill=rgb(INK))
        yy += row_h
    return (x, y, x + w, y + h)


def make_use_case():
    img = Image.new("RGB", (1800, 1200), "white")
    d = ImageDraw.Draw(img)
    title = "Modelo de casos de uso - Administracion de proyectos"
    bb = d.textbbox((0, 0), title, font=fnt(42, True))
    d.text(((1800 - (bb[2] - bb[0])) / 2, 35), title, font=fnt(42, True), fill=rgb(BLUE_DARK))

    boundary = (455, 135, 1615, 1080)
    d.rounded_rectangle(boundary, radius=18, outline=rgb(BLUE_DARK), width=4, fill=rgb("FBFDFF"))
    center(d, (455, 150, 1615, 210), "Sistema Administracion de proyectos", fnt(28, True), fill=rgb(BLUE_DARK))

    def actor(x, y, label):
        d.ellipse((x + 50, y, x + 105, y + 55), outline=rgb(INK), width=4)
        d.line((x + 78, y + 55, x + 78, y + 150), fill=rgb(INK), width=4)
        d.line((x + 25, y + 95, x + 130, y + 95), fill=rgb(INK), width=4)
        d.line((x + 78, y + 150, x + 35, y + 225), fill=rgb(INK), width=4)
        d.line((x + 78, y + 150, x + 120, y + 225), fill=rgb(INK), width=4)
        center(d, (x - 25, y + 238, x + 185, y + 310), label, fnt(22, True), fill=rgb(INK), max_width=190)

    actor(90, 170, "Direccion / Coordinacion")
    actor(90, 505, "Docente / Asesor")
    actor(90, 835, "Estudiante / Equipo")

    cases = {
        "catalogos": (560, 250, 900, 325, "Administrar periodos, usuarios, grupos y asignaturas"),
        "guias": (1035, 250, 1465, 325, "Crear y versionar guias integradoras"),
        "asignar": (1035, 395, 1465, 470, "Asignar docentes por apartado"),
        "reportes": (560, 395, 900, 470, "Consultar reportes y bitacora"),
        "ver_asignados": (560, 590, 930, 665, "Consultar apartados asignados"),
        "revisar": (1045, 590, 1465, 665, "Revisar entrega y evidencia"),
        "dictaminar": (1045, 735, 1465, 810, "Aprobar, rechazar o solicitar correccion"),
        "comentar": (560, 735, 930, 810, "Registrar comentarios de revision"),
        "consultar": (560, 930, 930, 1005, "Consultar guia y fechas"),
        "entregar": (1045, 930, 1465, 1005, "Entregar documentos y codigo"),
    }
    for x1, y1, x2, y2, text in cases.values():
        d.ellipse((x1, y1, x2, y2), outline=rgb(BLUE_DARK), width=3, fill=rgb(LIGHT_BLUE))
        center(d, (x1, y1, x2, y2), text, fnt(20), fill=rgb(INK), max_width=x2 - x1 - 45)

    line = rgb("93A0B1")
    for end in [(560, 287), (1035, 287), (1035, 432), (560, 432)]:
        poly_arrow(d, [(260, 315), (365, 315), (365, end[1]), end], fill=line, width=3)
    for end in [(560, 627), (1045, 627), (1045, 772), (560, 772)]:
        poly_arrow(d, [(260, 650), (380, 650), (380, end[1]), end], fill=line, width=3)
    for end in [(560, 967), (1045, 967)]:
        poly_arrow(d, [(260, 980), (390, 980), (390, end[1]), end], fill=line, width=3)

    d.text((980, 700), "<<include>>", font=fnt(18), fill=rgb(MUTED))
    poly_arrow(d, [(1255, 665), (1255, 735)], fill=rgb("B0B8C4"), width=3)
    d.text((720, 690), "<<include>>", font=fnt(18), fill=rgb(MUTED))
    poly_arrow(d, [(745, 665), (745, 735)], fill=rgb("B0B8C4"), width=3)
    for x1, y1, x2, y2, text in cases.values():
        d.ellipse((x1, y1, x2, y2), outline=rgb(BLUE_DARK), width=3, fill=rgb(LIGHT_BLUE))
        center(d, (x1, y1, x2, y2), text, fnt(20), fill=rgb(INK), max_width=x2 - x1 - 45)
    d.text((980, 700), "<<include>>", font=fnt(18), fill=rgb(MUTED))
    d.text((720, 690), "<<include>>", font=fnt(18), fill=rgb(MUTED))
    img.save(USE_CASE_IMG)


def make_er():
    img = Image.new("RGB", (1900, 1450), "white")
    d = ImageDraw.Draw(img)
    title = "Modelo entidad-relacion - Administracion de proyectos"
    bb = d.textbbox((0, 0), title, font=fnt(44, True))
    d.text(((1900 - (bb[2] - bb[0])) / 2, 35), title, font=fnt(44, True), fill=rgb(BLUE_DARK))

    specs = {
        "roles": (80, 165, 250, "roles", ["id PK", "nombre", "descripcion"], LIGHT_BLUE),
        "usuarios": (410, 155, 315, "usuarios", ["id PK", "nombre", "email", "matricula", "estatus"], LIGHT_BLUE),
        "periodos": (820, 155, 290, "periodos", ["id PK", "nombre", "fecha_inicio", "fecha_fin"], LIGHT_BLUE),
        "asignaturas": (1200, 155, 320, "asignaturas", ["id PK", "nombre", "clave", "cuatrimestre"], LIGHT_BLUE),
        "grupos": (1600, 155, 245, "grupos", ["id PK", "nombre", "carrera", "periodo_id FK"], LIGHT_BLUE),
        "guias": (110, 495, 340, "guias_integradoras", ["id PK", "asignatura_id FK", "periodo_id FK", "version", "estatus"], LIGHT_GREEN),
        "apartados": (560, 495, 360, "apartados_guia", ["id PK", "guia_id FK", "titulo", "fecha_entrega", "ponderacion"], LIGHT_GREEN),
        "equipos": (1450, 495, 330, "equipos", ["id PK", "grupo_id FK", "nombre", "lider_id FK"], LIGHT_GREEN),
        "proyectos": (1010, 505, 335, "proyectos", ["id PK", "guia_id FK", "equipo_id FK", "titulo", "estado"], LIGHT_GREEN),
        "integrantes": (1450, 785, 330, "equipo_integrantes", ["id PK", "equipo_id FK", "estudiante_id FK", "activo"], "FFF7E6"),
        "asignaciones": (535, 835, 390, "asignaciones_revision", ["id PK", "proyecto_id FK", "apartado_id FK", "docente_id FK", "activo"], "FFF7E6"),
        "entregas": (1010, 835, 360, "entregas", ["id PK", "proyecto_id FK", "apartado_id FK", "version", "estado"], "FFF7E6"),
        "revisiones": (535, 1140, 350, "revisiones", ["id PK", "entrega_id FK", "revisor_id FK", "resultado", "calificacion"], "FFF7E6"),
        "comentarios": (105, 1140, 350, "comentarios_revision", ["id PK", "revision_id FK", "autor_id FK", "comentario"], "FFF7E6"),
        "archivos": (1010, 1140, 360, "archivos_entrega", ["id PK", "entrega_id FK", "tipo", "ruta", "version"], "FFF7E6"),
        "codigo": (1450, 1140, 330, "productos_codigo", ["id PK", "proyecto_id FK", "repositorio_url", "archivo", "version"], "FFF7E6"),
    }

    boxes = {}
    for key, (x, y, w, title_text, attrs, fill) in specs.items():
        boxes[key] = (x, y, x + w, y + 46 + 30 * len(attrs))

    def mid(name, side):
        x1, y1, x2, y2 = boxes[name]
        if side == "r":
            return (x2, (y1 + y2) // 2)
        if side == "l":
            return (x1, (y1 + y2) // 2)
        if side == "b":
            return ((x1 + x2) // 2, y2)
        return ((x1 + x2) // 2, y1)

    links = [
        ("roles", "usuarios", "r", "l"),
        ("periodos", "guias", "b", "t"),
        ("asignaturas", "guias", "b", "t"),
        ("grupos", "equipos", "b", "t"),
        ("guias", "apartados", "r", "l"),
        ("guias", "proyectos", "r", "l"),
        ("equipos", "proyectos", "l", "r"),
        ("equipos", "integrantes", "b", "t"),
        ("apartados", "asignaciones", "b", "t"),
        ("proyectos", "asignaciones", "b", "t"),
        ("apartados", "entregas", "b", "t"),
        ("proyectos", "entregas", "b", "t"),
        ("entregas", "revisiones", "l", "r"),
        ("revisiones", "comentarios", "l", "r"),
        ("entregas", "archivos", "b", "t"),
        ("proyectos", "codigo", "r", "l"),
        ("usuarios", "integrantes", "r", "l"),
        ("usuarios", "asignaciones", "r", "l"),
    ]
    for a, b, sa, sb in links:
        arrow(d, mid(a, sa), mid(b, sb), fill=rgb("A8B3C2"), width=3)

    for key, (x, y, w, title_text, attrs, fill) in specs.items():
        entity(d, x, y, w, title_text, attrs, fill)

    img.save(ER_IMG)


def make_relational_summary():
    img = Image.new("RGB", (2300, 1750), "white")
    d = ImageDraw.Draw(img)
    title = "Modelo relacional completo - Administracion de proyectos"
    bb = d.textbbox((0, 0), title, font=fnt(42, True))
    d.text(((2300 - (bb[2] - bb[0])) / 2, 35), title, font=fnt(42, True), fill=rgb(BLUE_DARK))

    specs = {
        "roles": (70, 150, 280, "roles", ["id PK", "nombre", "descripcion"], LIGHT_BLUE),
        "usuarios": (430, 140, 335, "usuarios", ["id PK", "nombre", "email", "password", "matricula", "estatus"], LIGHT_BLUE),
        "usuario_roles": (845, 150, 330, "usuario_roles", ["id PK", "usuario_id FK", "rol_id FK"], LIGHT_BLUE),
        "periodos": (1260, 150, 310, "periodos", ["id PK", "nombre", "fecha_inicio", "fecha_fin", "estatus"], LIGHT_BLUE),
        "asignaturas": (1660, 150, 345, "asignaturas", ["id PK", "nombre", "clave", "cuatrimestre", "estatus"], LIGHT_BLUE),
        "grupos": (70, 520, 300, "grupos", ["id PK", "periodo_id FK", "nombre", "carrera", "cuatrimestre"], LIGHT_GREEN),
        "guias": (470, 520, 360, "guias_integradoras", ["id PK", "asignatura_id FK", "periodo_id FK", "nombre", "version", "estatus"], LIGHT_GREEN),
        "apartados": (930, 520, 405, "apartados_guia", ["id PK", "guia_id FK", "orden", "titulo", "fecha_entrega", "ponderacion", "requiere_codigo"], LIGHT_GREEN),
        "equipos": (1450, 520, 330, "equipos", ["id PK", "grupo_id FK", "lider_id FK", "nombre", "estatus"], LIGHT_GREEN),
        "proyectos": (1885, 520, 350, "proyectos", ["id PK", "guia_id FK", "equipo_id FK", "titulo", "descripcion", "estado"], LIGHT_GREEN),
        "equipo_integrantes": (70, 940, 360, "equipo_integrantes", ["id PK", "equipo_id FK", "estudiante_id FK", "activo"], "FFF7E6"),
        "asignaciones": (535, 940, 405, "asignaciones_revision", ["id PK", "proyecto_id FK", "apartado_guia_id FK", "docente_id FK", "tipo_revisor", "fecha_inicio", "fecha_fin", "activo"], "FFF7E6"),
        "historial": (1045, 940, 405, "historial_asignaciones", ["id PK", "asignacion_revision_id FK", "docente_anterior_id FK", "docente_nuevo_id FK", "motivo", "fecha"], "FFF7E6"),
        "entregas": (1560, 940, 370, "entregas", ["id PK", "proyecto_id FK", "apartado_guia_id FK", "equipo_id FK", "version", "estado_entrega", "fecha_entrega"], "FFF7E6"),
        "archivos": (70, 1370, 365, "archivos_entrega", ["id PK", "entrega_id FK", "nombre_original", "ruta", "tipo_archivo", "tamano"], "F2F4F7"),
        "codigo": (545, 1370, 390, "productos_codigo", ["id PK", "proyecto_id FK", "entrega_id FK", "repositorio_url", "archivo_fuente", "version"], "F2F4F7"),
        "revisiones": (1045, 1370, 395, "revisiones", ["id PK", "entrega_id FK", "asignacion_revision_id FK", "revisor_id FK", "resultado", "calificacion"], "F2F4F7"),
        "comentarios": (1535, 1370, 390, "comentarios_revision", ["id PK", "revision_id FK", "autor_id FK", "comentario", "visible_estudiante"], "F2F4F7"),
        "notificaciones": (1985, 940, 285, "notificaciones", ["id PK", "usuario_id FK", "titulo", "mensaje", "leido", "fecha"], "F2F4F7"),
        "bitacora": (1985, 1370, 285, "bitacora_actividades", ["id PK", "usuario_id FK", "modulo", "accion", "descripcion", "fecha"], "F2F4F7"),
    }

    boxes = {}
    for key, (x, y, w, title_text, attrs, fill) in specs.items():
        boxes[key] = (x, y, x + w, y + 46 + 30 * len(attrs))

    def port(name, side, offset=0):
        x1, y1, x2, y2 = boxes[name]
        if side == "r":
            return (x2, (y1 + y2) // 2 + offset)
        if side == "l":
            return (x1, (y1 + y2) // 2 + offset)
        if side == "b":
            return ((x1 + x2) // 2 + offset, y2)
        return ((x1 + x2) // 2 + offset, y1)

    line = rgb("A8B3C2")
    routes = [
        [port("roles", "r"), (390, 220), port("usuario_roles", "l", -20)],
        [port("usuarios", "r"), (805, 260), port("usuario_roles", "l", 25)],
        [port("periodos", "b"), (1415, 460), (220, 460), port("grupos", "t")],
        [port("periodos", "b", 70), (1490, 430), (650, 430), port("guias", "t")],
        [port("asignaturas", "b"), (1830, 455), (660, 455), port("guias", "t", 70)],
        [port("grupos", "r"), (420, 650), (420, 730), (1450, 730), port("equipos", "l")],
        [port("guias", "r"), (880, 650), port("apartados", "l")],
        [port("guias", "r", 40), (1660, 830), (2060, 830), port("proyectos", "b", -70)],
        [port("apartados", "b", -60), (1130, 900), port("asignaciones", "t")],
        [port("apartados", "b", 80), (1180, 895), (1745, 895), port("entregas", "t", -70)],
        [port("equipos", "b"), (1615, 900), (250, 900), port("equipo_integrantes", "t")],
        [port("equipos", "r"), (1830, 650), port("proyectos", "l")],
        [port("proyectos", "b", -65), (1985, 895), (735, 895), port("asignaciones", "t", 80)],
        [port("proyectos", "b", 30), (2065, 895), (1745, 895), port("entregas", "t", 70)],
        [port("proyectos", "b", 85), (2135, 1290), (740, 1290), port("codigo", "t")],
        [port("asignaciones", "r"), (995, 1065), port("historial", "l")],
        [port("asignaciones", "b"), (740, 1325), (1230, 1325), port("revisiones", "t", -70)],
        [port("entregas", "b", -80), (1730, 1325), (250, 1325), port("archivos", "t")],
        [port("entregas", "b"), (1750, 1325), (740, 1325), port("codigo", "t", 80)],
        [port("entregas", "b", 80), (1780, 1325), (1240, 1325), port("revisiones", "t", 70)],
        [port("revisiones", "r"), (1490, 1490), port("comentarios", "l")],
        [port("usuarios", "b", -90), (520, 900), (247, 900), port("equipo_integrantes", "t", -70)],
        [port("usuarios", "b"), (590, 900), (738, 900), port("asignaciones", "t", -80)],
        [port("usuarios", "b", 85), (690, 1285), (1240, 1285), port("revisiones", "t")],
        [port("usuarios", "r"), (780, 350), (2125, 350), port("notificaciones", "t")],
        [port("usuarios", "r", 35), (800, 380), (2125, 380), (2125, 1370)],
    ]
    for route in routes:
        poly_arrow(d, route, fill=line, width=3)

    for key, (x, y, w, title_text, attrs, fill) in specs.items():
        entity(d, x, y, w, title_text, attrs, fill)

    img.save(REL_IMG)


def make_sitemap():
    img = Image.new("RGB", (1700, 1150), "white")
    d = ImageDraw.Draw(img)
    title = "Mapa de sitio - Administracion de proyectos"
    bb = d.textbbox((0, 0), title, font=fnt(42, True))
    d.text(((1700 - (bb[2] - bb[0])) / 2, 35), title, font=fnt(42, True), fill=rgb(BLUE_DARK))
    root = (675, 120, 1025, 195)
    d.rounded_rectangle(root, radius=14, fill=rgb(BLUE), outline=rgb(BLUE_DARK), width=3)
    center(d, root, "Inicio / Login", fnt(26, True), fill="white")
    columns = [
        ("Direccion / Coordinacion", ["Dashboard", "Periodos", "Usuarios y roles", "Asignaturas y grupos", "Guias integradoras", "Asignar revisores", "Reportes", "Bitacora"], 70, LIGHT_BLUE),
        ("Docente / Asesor", ["Dashboard de revision", "Apartados asignados", "Entregas por revisar", "Comentarios", "Historial de revisiones"], 500, LIGHT_GREEN),
        ("Estudiante / Equipo", ["Dashboard del proyecto", "Guia integradora", "Entregar avance", "Subir documentos", "Subir codigo", "Ver comentarios", "Reentregar correcciones"], 930, "FFF7E6"),
        ("Soporte del sistema", ["Notificaciones", "Perfil", "Ayuda", "Cerrar sesion"], 1360, "F2F4F7"),
    ]
    for title, items, x, fill in columns:
        y = 300
        box = (x, y, x + 290, y + 70)
        arrow(d, ((root[0] + root[2])//2, root[3]), (x + 145, y), fill=rgb("8A97A8"), width=3)
        d.rounded_rectangle(box, radius=12, fill=rgb(fill), outline=rgb(BLUE_DARK), width=3)
        center(d, box, title, fnt(21, True), fill=rgb(INK), max_width=250)
        yy = y + 105
        for item in items:
            item_box = (x + 18, yy, x + 272, yy + 50)
            d.rounded_rectangle(item_box, radius=10, fill="white", outline=rgb(BORDER), width=2)
            center(d, item_box, item, fnt(19), fill=rgb(INK), max_width=220)
            arrow(d, (x + 145, yy - 35), (x + 145, yy), fill=rgb("B0B8C4"), width=2)
            yy += 72
    img.save(SITE_IMG)


def set_font(run, size=10.5, bold=False, color=INK):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)
        tc_borders.append(tag)


def add_p(doc, text="", size=10.5, bold=False, color=INK, align=None, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    if align:
        p.alignment = align
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    return p


def add_h(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(15 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, size=16 if level == 1 else 13, bold=True, color=BLUE if level == 1 else GREEN_DARK)
    return p


def add_table(doc, headers, rows, widths=None, font_size=8.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        set_font(r, size=font_size, bold=True, color="FFFFFF")
        shade(cell, BLUE)
        borders(cell)
    for ridx, row in enumerate(rows, start=1):
        cells = table.add_row().cells
        for i, val in enumerate(row[:len(headers)]):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(val))
            set_font(r, size=font_size, color=INK)
            shade(cells[i], "FFFFFF" if ridx % 2 else LIGHT_GRAY)
            borders(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths[:len(row.cells)]):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_font(r, size=9, color=MUTED)
    r.italic = True


def build_doc():
    make_logo()
    make_use_case()
    make_er()
    make_relational_summary()
    make_sitemap()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles["Normal"]
    styles.font.name = "Arial"
    styles.font.size = Pt(10.5)

    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Inches(4.4))

    add_p(doc, "Administracion de proyectos", size=24, bold=True, color=BLUE_DARK, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    add_p(doc, "Propuesta documental para sistema de evaluacion de proyectos estudiantiles", size=14, bold=True, color=GREEN_DARK, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    add_p(doc, "Universidad Tecnologica del Valle del Mezquital", size=11, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    add_p(doc, "Grupo: 9°B", size=11, bold=True, color=BLUE_DARK, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)

    add_h(doc, "Equipo de trabajo", 1)
    add_table(doc, ["Integrante", "Rol en el proyecto"], [
        ["Miguel Espinal Botho", "Tester", "9°B"],
        ["Juan Pablo Barrera Santiago", "Programador", "9°B"],
        ["David Martinez Polvadera", "Programador", "9°B"],
        ["Said Castillo Aguilar", "Programador de base de datos", "9°B"],
        ["Nayeli Ortiz Dimas", "Diseñadora", "9°B"],
        ["Leslie Mayte Vicentes Acosta", "Scrum Master", "9°B"],
        ["Flor Mendoza Hernández", "Product Owner", "9°B"],
    ], widths=[2.55, 2.65, 0.75], font_size=8.8)

    add_h(doc, "1. Proposito del sistema", 1)
    add_p(doc, "El sistema Administracion de proyectos se propone como una plataforma academica para controlar la entrega, revision, correccion y calificacion de proyectos elaborados por estudiantes. La base de trabajo sera la guia integradora definida por la direccion o coordinacion academica, incluyendo contenido, fechas de entrega, responsables de revision y evidencias solicitadas.")
    add_p(doc, "La plataforma debe permitir que cada docente o asesor visualice unicamente los apartados que le corresponde revisar. Esto evita que el calificador tenga acceso operativo a secciones que no forman parte de su responsabilidad y permite distribuir la evaluacion entre varios docentes cuando la guia integradora involucra asignaturas distintas.")

    add_h(doc, "2. Identidad visual propuesta", 1)
    add_p(doc, "La identidad visual toma como referencia el sitio oficial de UTVM y las pantallas compartidas: uso dominante de azul institucional, verde como color academico/confirmacion, acentos turquesa para informacion y amarillo para avisos o acciones secundarias.")
    add_table(doc, ["Uso", "Color", "Hexadecimal", "Aplicacion sugerida"], [
        ["Primario institucional", "Azul UTVM", "#15529A", "Barra superior, encabezados principales, botones primarios."],
        ["Primario oscuro", "Azul profundo", "#0D376D", "Titulos, texto sobre fondos claros, pie de pagina."],
        ["Secundario", "Verde UTVM", "#21A366", "Estados aprobados, indicadores positivos, acentos."],
        ["Secundario oscuro", "Verde academico", "#0F7D47", "Subtitulos y etiquetas de modulo."],
        ["Informativo", "Turquesa", "#1EA7D7", "Avisos informativos, chips de estado, enlaces."],
        ["Advertencia", "Amarillo", "#F3B21A", "Alertas, pendientes, confirmar acciones."],
        ["Neutro", "Gris claro", "#F5F7FA", "Fondos de paneles, tablas y formularios."],
    ], widths=[1.25, 1.2, 1.0, 3.2])

    add_h(doc, "3. Reglas funcionales base", 1)
    rules = [
        ["RF-01", "La direccion o coordinacion registra periodos, asignaturas, grupos, usuarios, docentes y guias integradoras."],
        ["RF-02", "Cada guia integradora contiene apartados, fechas de entrega, ponderaciones, tipo de evidencia y docente/asesor responsable."],
        ["RF-03", "Un proyecto puede tener varios docentes revisores, cada uno asignado a uno o varios apartados."],
        ["RF-04", "Los docentes y asesores solo visualizan los apartados que tienen asignados para revision."],
        ["RF-05", "Las asignaciones de docentes pueden cambiar durante el periodo sin eliminar el historial anterior."],
        ["RF-06", "Las entregas pueden incluir documentos, evidencias, comentarios, enlaces a repositorios y archivos de codigo."],
        ["RF-07", "Toda revision debe registrar resultado: pendiente, en revision, aprobado, rechazado o corregir."],
        ["RF-08", "Cuando una revision sea rechazada o marcada para corregir, el comentario del docente sera obligatorio."],
        ["RF-09", "El sistema debe generar reportes de avance, cumplimiento, retrasos, docentes asignados y estado de entregas."],
    ]
    add_table(doc, ["Clave", "Regla"], rules, widths=[0.75, 5.75], font_size=9)

    add_h(doc, "4. Modelo de casos de uso", 1)
    add_p(doc, "El modelo de casos de uso describe las funciones principales por actor. La direccion administra la configuracion academica, el docente o asesor revisa solamente apartados asignados, y el estudiante entrega avances, documentos y codigo.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(USE_CASE_IMG), width=Inches(7.0))
    add_caption(doc, "Figura 1. Modelo de casos de uso propuesto.")

    add_h(doc, "5. Modelo entidad-relacion", 1)
    add_p(doc, "El modelo entidad-relacion define las entidades principales que soportan la operacion academica. La entidad central para resolver los cambios de docentes y la revision parcial por apartado es asignaciones_revision.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ER_IMG), width=Inches(7.15))
    add_caption(doc, "Figura 2. Modelo entidad-relacion propuesto.")

    add_h(doc, "6. Modelo relacional", 1)
    add_p(doc, "El modelo relacional completo presenta las tablas principales del sistema, sus llaves primarias y llaves foraneas. Este modelo evita manejar docentes o asesores como datos fijos del proyecto, ya que las asignaciones de revision permiten cambios durante el periodo.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(REL_IMG), width=Inches(7.2))
    add_caption(doc, "Figura 3. Modelo relacional completo propuesto.")

    add_table(doc, ["Tabla", "Campos principales", "Llaves y relacion"], [
        ["roles", "id, nombre, descripcion", "PK id. Relacion N:M con usuarios mediante usuario_roles."],
        ["usuarios", "id, nombre, email, password, matricula, estatus", "PK id. Participa como estudiante, docente, asesor, coordinador o administrador."],
        ["usuario_roles", "id, usuario_id, rol_id", "FK usuario_id, rol_id. Permite que un docente tambien sea asesor o coordinador."],
        ["periodos", "id, nombre, fecha_inicio, fecha_fin, estatus", "PK id. Relacion con grupos y guias."],
        ["asignaturas", "id, nombre, clave, cuatrimestre, estatus", "PK id. Relacion con guias integradoras."],
        ["grupos", "id, nombre, carrera, cuatrimestre, periodo_id", "FK periodo_id. Agrupa estudiantes por periodo academico."],
        ["guias_integradoras", "id, asignatura_id, periodo_id, nombre, version, estatus", "FK asignatura_id, periodo_id. Base formal del proyecto."],
        ["apartados_guia", "id, guia_id, orden, titulo, descripcion, fecha_entrega, ponderacion, requiere_documento, requiere_codigo", "FK guia_id. Define que debe entregar el estudiante."],
        ["equipos", "id, grupo_id, nombre, lider_id, estatus", "FK grupo_id, lider_id. Equipo de estudiantes."],
        ["equipo_integrantes", "id, equipo_id, estudiante_id, activo", "FK equipo_id, estudiante_id. Historial de integrantes por equipo."],
        ["proyectos", "id, guia_id, equipo_id, titulo, descripcion, estado", "FK guia_id, equipo_id. Proyecto trabajado por un equipo."],
        ["asignaciones_revision", "id, proyecto_id, apartado_guia_id, docente_id, tipo_revisor, fecha_inicio, fecha_fin, activo", "FK proyecto_id, apartado_guia_id, docente_id. Controla quien revisa cada parte."],
        ["historial_asignaciones", "id, asignacion_revision_id, docente_anterior_id, docente_nuevo_id, motivo, fecha", "FK asignacion_revision_id. Conserva cambios de docente o asesor."],
        ["entregas", "id, proyecto_id, apartado_guia_id, equipo_id, version, estado_entrega, fecha_entrega", "FK proyecto_id, apartado_guia_id, equipo_id. Registra cada envio del estudiante."],
        ["archivos_entrega", "id, entrega_id, nombre_original, ruta, tipo_archivo, tamano", "FK entrega_id. Documentos y evidencias adjuntas."],
        ["productos_codigo", "id, proyecto_id, entrega_id, repositorio_url, archivo_fuente, version, observaciones", "FK proyecto_id, entrega_id. Codigo fuente o repositorio."],
        ["revisiones", "id, entrega_id, asignacion_revision_id, revisor_id, resultado, calificacion, comentario_general", "FK entrega_id, asignacion_revision_id, revisor_id. Resultado: aprobado, rechazado o corregir."],
        ["comentarios_revision", "id, revision_id, autor_id, comentario, visible_estudiante", "FK revision_id, autor_id. Observaciones para retroalimentacion."],
        ["notificaciones", "id, usuario_id, titulo, mensaje, leido, fecha", "FK usuario_id. Avisos de entregas, revisiones y correcciones."],
        ["bitacora_actividades", "id, usuario_id, modulo, accion, descripcion, fecha", "FK usuario_id. Auditoria del sistema."],
    ], widths=[1.45, 2.55, 2.5], font_size=7.4)

    add_h(doc, "7. Estados y permisos de revision", 1)
    add_table(doc, ["Estado", "Significado", "Accion esperada"], [
        ["Pendiente", "El apartado aun no ha sido entregado o no tiene revision iniciada.", "Esperar entrega o asignar responsable."],
        ["En revision", "El docente/asesor esta evaluando la entrega.", "Registrar observaciones o resultado."],
        ["Aprobado", "La entrega cumple con lo solicitado.", "Cerrar revision y acumular avance/calificacion."],
        ["Rechazado", "La entrega no cumple requisitos minimos.", "Comentario obligatorio y posible nueva entrega."],
        ["Corregir", "La entrega tiene observaciones solucionables.", "Comentario obligatorio y reentrega del estudiante."],
        ["Fuera de tiempo", "La entrega fue enviada despues de la fecha limite.", "Registrar entrega y aplicar regla academica definida."],
    ], widths=[1.25, 2.8, 2.45])

    add_h(doc, "8. Mapa de sitio", 1)
    add_p(doc, "El mapa de sitio separa las pantallas por rol. Esto permite que el estudiante trabaje sobre sus entregas, el docente revise solo sus apartados asignados y direccion mantenga el control general del proceso.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(SITE_IMG), width=Inches(7.0))
    add_caption(doc, "Figura 4. Mapa de sitio propuesto por rol.")

    add_h(doc, "9. Criterios de aprobacion documental", 1)
    add_table(doc, ["Elemento", "Criterio de aceptacion"], [
        ["Identidad visual", "Uso de logo institucional UTVM, paleta azul/verde y componentes formales."],
        ["Modelo ER", "Incluye guias, apartados, proyectos, equipos, entregas, revisiones, comentarios y cambios de asignacion."],
        ["Modelo relacional", "Define tablas, campos principales, PK/FK y objetivo de cada relacion."],
        ["Mapa de sitio", "Diferencia claramente direccion/coordinacion, docente/asesor y estudiante/equipo."],
        ["Reglas de revision", "Limita visibilidad del docente a apartados asignados y registra estados de aprobacion/correccion."],
    ], widths=[1.7, 4.8])

    add_p(doc, "Fuente visual consultada: sitio oficial de la Universidad Tecnologica del Valle del Mezquital, https://www.utvm.edu.mx/.", size=9, color=MUTED)
    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
